"""Explore official early-vote sources and their archived daily reports (runs in GitHub Actions; the sandbox can't reach them).

    python3 pipeline/ev_sources.py explore OUTDIR

Writes small text files describing what each source offers: Wayback Machine snapshot lists for the Pennsylvania and
Florida daily reports in 2022 and 2024, a sample of each, and the link lists of the Nevada, Colorado, Iowa and
Georgia election-office pages. Standard library plus openpyxl (installed by the workflow).
"""
import io
import json
import os
import re
import sys
import urllib.parse
import urllib.request

UA = {'User-Agent': 'midterm-map early-vote research (github.com/DanCunningham22/2026-Midterm)'}


def get(url, timeout=120, binary=False):
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=timeout) as r:
        data = r.read()
    return data if binary else data.decode('utf-8', errors='replace')


def cdx(url, frm, to, limit=400):
    q = ('https://web.archive.org/cdx/search/cdx?url=' + urllib.parse.quote(url, safe='') +
         f'&output=json&fl=timestamp,statuscode,mimetype,length&from={frm}&to={to}&limit={limit}&filter=statuscode:200')
    try:
        rows = json.loads(get(q))
    except Exception as e:
        return [], str(e)
    return rows[1:], ''


def wayback(ts, url, binary=False):
    return get(f'https://web.archive.org/web/{ts}id_/{url}', binary=binary)


def links(html, base):
    out = []
    for m in re.finditer(r'<a[^>]+href=["\']([^"\']+)["\'][^>]*>(.*?)</a>', html, re.I | re.S):
        href, text = m.group(1), re.sub(r'<[^>]+>', '', m.group(2)).strip()
        out.append((urllib.parse.urljoin(base, href), re.sub(r'\s+', ' ', text)[:80]))
    return out


def xlsx_preview(data, rows=25):
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f'--- sheet {ws.title} (max_row {ws.max_row}, max_col {ws.max_column})')
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= rows:
                break
            out.append(' | '.join('' if v is None else str(v) for v in row))
    return '\n'.join(out)


def text_preview(html, n=6000):
    t = re.sub(r'<script.*?</script>|<style.*?</style>', ' ', html, flags=re.S | re.I)
    t = re.sub(r'<[^>]+>', ' ', t)
    t = re.sub(r'\s+', ' ', t)
    return t[:n]


def explore(outdir):
    os.makedirs(outdir, exist_ok=True)
    log = []

    def save(name, text):
        with open(os.path.join(outdir, name), 'w', encoding='utf-8') as f:
            f.write(text)
        log.append(f'wrote {name} ({len(text)} chars)')

    # Pennsylvania: the daily mail ballot report (xlsx), archived by the Wayback Machine
    for yr in ('2022', '2024', '2026'):
        url = f'https://www.pavoterservices.pa.gov/{yr}%20General%20Daily%20Mail%20Ballot%20Report.xlsx'
        rows, err = cdx(url, f'{yr}0901', f'{yr}1110')
        save(f'pa_{yr}_cdx.txt', (err or '') + '\n' + '\n'.join(' '.join(r) for r in rows))
        pick = rows[len(rows) // 2] if rows else None
        if pick:
            try:
                data = wayback(pick[0], url, binary=True)
                save(f'pa_{yr}_sample_{pick[0]}.txt', xlsx_preview(data))
            except Exception as e:
                save(f'pa_{yr}_sample_error.txt', str(e))
        if yr == '2026':
            try:
                data = get(url, binary=True)
                save('pa_2026_live.txt', xlsx_preview(data))
            except Exception as e:
                save('pa_2026_live_error.txt', str(e))

    # Florida: the Division of Elections' vote-by-mail and early-voting report page (statewide totals by party)
    fl = 'https://dos.fl.gov/elections/data-statistics/elections-data/vote-by-mail-and-early-voting-reports/'
    for yr in ('2022', '2024'):
        rows, err = cdx(fl, f'{yr}1001', f'{yr}1110')
        save(f'fl_{yr}_cdx.txt', (err or '') + '\n' + '\n'.join(' '.join(r) for r in rows))
        if rows:
            pick = rows[len(rows) // 2]
            try:
                html = wayback(pick[0], fl)
                save(f'fl_{yr}_sample_{pick[0]}.txt', text_preview(html, 12000) + '\n\nLINKS\n' + '\n'.join(f'{u} :: {t}' for u, t in links(html, fl)[:150]))
            except Exception as e:
                save(f'fl_{yr}_sample_error.txt', str(e))
    try:
        html = get(fl)
        save('fl_2026_live.txt', text_preview(html, 12000) + '\n\nLINKS\n' + '\n'.join(f'{u} :: {t}' for u, t in links(html, fl)[:150]))
    except Exception as e:
        save('fl_2026_live_error.txt', str(e))

    # Nevada, Colorado, Iowa, Georgia, Arizona: what the election offices link to
    pages = {
        'nv_stats': 'https://www.nvsos.gov/sos/elections/election-information/election-statistics',
        'nv_elections': 'https://www.nvsos.gov/sos/elections',
        'co_returns': 'https://www.sos.state.co.us/pubs/elections/VoteByMail/ballotReturns.html',
        'co_elections': 'https://www.sos.state.co.us/pubs/elections/',
        'ia_absentee': 'https://sos.iowa.gov/elections/electioninfo/absenteestats.html',
        'ia_elections': 'https://sos.iowa.gov/elections/',
        'ga_absentee': 'https://sos.ga.gov/page/voter-absentee-files',
        'ga_elections': 'https://sos.ga.gov/elections-division-georgia-secretary-states-office',
        'az_sos': 'https://azsos.gov/elections',
    }
    for k, u in pages.items():
        try:
            html = get(u)
            ls = [f'{a} :: {t}' for a, t in links(html, u) if re.search(r'turnout|absentee|early|mail|ballot|return|statistic', a + ' ' + t, re.I)]
            save(f'{k}.txt', u + '\n' + text_preview(html, 3000) + '\n\nLINKS\n' + '\n'.join(ls[:200]))
        except Exception as e:
            save(f'{k}_error.txt', u + '\n' + str(e))
    save('_log.txt', '\n'.join(log))


def explore2(outdir):
    """Second pass: the UF Election Lab's data endpoints, Colorado's daily spreadsheets, Florida's statistics page,
    the 2022 RPubs tracker, Pennsylvania's archived report filenames, Nevada with a browser user agent."""
    import datetime
    os.makedirs(outdir, exist_ok=True)
    log = []

    def save(name, text):
        with open(os.path.join(outdir, name), 'w', encoding='utf-8') as f:
            f.write(text)
        log.append(f'wrote {name} ({len(text)} chars)')

    def tryget(url, **kw):
        try:
            return get(url, **kw), ''
        except Exception as e:
            return '', str(e)

    # 1. UF Election Lab: where do the pages get their numbers?
    uf_pages = {
        'uf26_nat': 'https://election.lab.ufl.edu/early-vote/2026-early-voting/',
        'uf26_pa': 'https://election.lab.ufl.edu/early-vote/2026-early-voting/2026-general-election-early-vote-pennsylvania/',
        'uf24_pa': 'https://election.lab.ufl.edu/early-vote/2024-early-voting/2024-general-election-early-vote-pennsylvania/',
    }
    scripts = []
    for k, u in uf_pages.items():
        html, err = tryget(u)
        if err:
            save(f'{k}_error.txt', err)
            continue
        srcs = re.findall(r'<script[^>]+src=["\']([^"\']+)["\']', html, re.I)
        hits = re.findall(r'["\']([^"\']*(?:wp-json|/api/|ajax|\.json|rest_route|admin-ajax|early[-_]?vote)[^"\']*)["\']', html, re.I)
        inline = re.findall(r'<script(?![^>]*src)[^>]*>(.*?)</script>', html, re.S | re.I)
        inline_hits = [s[:1500] for s in inline if re.search(r'fetch\(|ajax|wp-json|\.json|XMLHttpRequest|api', s, re.I)]
        save(f'{k}.txt', u + '\n\nSCRIPTS\n' + '\n'.join(srcs) + '\n\nURL HITS\n' + '\n'.join(sorted(set(hits))[:200]) +
             '\n\nINLINE SCRIPTS WITH FETCH/AJAX\n' + '\n\n=====\n'.join(inline_hits[:12]) + '\n\nTEXT\n' + text_preview(html, 4000))
        scripts += [s for s in srcs if 'election.lab.ufl.edu' in s and 'wp-includes' not in s]
    seen = set()
    for s in scripts:
        if s in seen:
            continue
        seen.add(s)
        js, err = tryget(s)
        if err:
            continue
        hits = re.findall(r'["\'`]([^"\'`\s]*(?:wp-json|/api/|ajax|\.json|\.csv|rest_route|admin-ajax|early)[^"\'`\s]*)["\'`]', js, re.I)
        if hits:
            name = re.sub(r'[^A-Za-z0-9]+', '_', s.split('election.lab.ufl.edu')[-1])[:80]
            save(f'ufjs_{name}.txt', s + '\n' + '\n'.join(sorted(set(hits))[:300]))
    for k, u in {'uf_wpjson': 'https://election.lab.ufl.edu/wp-json/', 'uf_wpjson_ev': 'https://election.lab.ufl.edu/wp-json/early-vote/v1/',
                 'uf_wpjson_ufev': 'https://election.lab.ufl.edu/wp-json/ufelectionlab/v1/'}.items():
        t, err = tryget(u)
        if err:
            save(f'{k}_error.txt', err)
        else:
            routes = re.findall(r'"(/[^"]*)":\{', t)
            save(f'{k}.txt', u + '\n' + '\n'.join(r for r in routes if re.search(r'early|vote|ev', r, re.I))[:20000] + '\n\nHEAD\n' + t[:3000])

    # 2. Colorado: the daily ballot-return press releases and spreadsheets (2022, 2024, 2026)
    co = []
    for yr, start, end in (('2022', datetime.date(2022, 10, 20), datetime.date(2022, 11, 9)),
                           ('2024', datetime.date(2024, 10, 14), datetime.date(2024, 11, 6)),
                           ('2026', datetime.date(2026, 10, 1), datetime.date.today())):
        d = start
        while d <= end:
            tag = d.strftime('%Y%m%d')
            u = f'https://www.coloradosos.gov/pubs/newsRoom/pressReleases/{yr}/PR{tag}BallotReturns.html'
            html, err = tryget(u, timeout=40)
            if not err:
                xl = re.findall(r'href=["\']([^"\']+\.xlsx)["\']', html, re.I)
                m = re.search(r'([\d,]{5,})\s+ballots', html)
                co.append(f'{tag} OK total={m.group(1) if m else "?"} xlsx={xl}')
                if xl and (len(co) <= 2 or d == end or d.strftime('%m%d') in ('1101', '1104', '1105', '1107')):
                    try:
                        data = get(urllib.parse.urljoin(u, xl[0]), binary=True)
                        save(f'co_{tag}_xlsx.txt', xlsx_preview(data, rows=40))
                    except Exception as e:
                        co.append(f'   xlsx error {e}')
            else:
                co.append(f'{tag} -- {err[:60]}')
            d += datetime.timedelta(days=1)
    save('co_releases.txt', '\n'.join(co))

    # 3. Florida: the statistics page (live) and its archived copies
    for k, u in {'fl_stats': 'https://countyfilesvbm-ev.floridados.gov/VoteByMailEarlyVotingReports/PublicStats',
                 'fl_vbm': 'https://countyfilesvbm-ev.floridados.gov/VoteByMailEarlyVotingReports/VoteByMailHome'}.items():
        html, err = tryget(u)
        if err:
            save(f'{k}_error.txt', err)
        else:
            forms = re.findall(r'<(select|form|input)[^>]*>', html, re.I)
            save(f'{k}.txt', u + '\n' + text_preview(html, 15000) + '\n\nFORMS\n' + '\n'.join(forms[:60]) + '\n\nLINKS\n' +
                 '\n'.join(f'{a} :: {t}' for a, t in links(html, u)[:200]))
    for host in ('countyballotfiles.floridados.gov', 'countyfilesvbm-ev.floridados.gov'):
        for yr in ('2022', '2024'):
            rows, err = cdx(f'{host}/VoteByMailEarlyVotingReports/PublicStats*', f'{yr}1001', f'{yr}1110')
            save(f'fl_cdx_{host.split(".")[0]}_{yr}.txt', (err or '') + '\n' + '\n'.join(' '.join(r) for r in rows))
            if rows:
                pick = rows[len(rows) // 2]
                try:
                    html = wayback(pick[0], f'https://{host}/VoteByMailEarlyVotingReports/PublicStats')
                    save(f'fl_sample_{yr}_{pick[0]}.txt', text_preview(html, 15000))
                except Exception as e:
                    save(f'fl_sample_{yr}_error.txt', str(e))

    # 4. The 2022 tracker on RPubs (static HTML, so archived copies carry the numbers)
    rp = 'https://rpubs.com/ElectProject/early_vote_2022'
    html, err = tryget(rp)
    if err:
        save('rpubs_error.txt', err)
    else:
        frames = re.findall(r'<iframe[^>]+src=["\']([^"\']+)["\']', html, re.I)
        save('rpubs.txt', rp + '\nIFRAMES\n' + '\n'.join(frames) + '\n\nTEXT\n' + text_preview(html, 3000))
        for fr in frames[:1]:
            fr = urllib.parse.urljoin(rp, fr)
            inner, err2 = tryget(fr)
            save('rpubs_inner.txt', fr + '\n' + (err2 or text_preview(inner, 20000)))
            rows, err3 = cdx(fr.replace('https://', ''), '20221001', '20221110')
            save('rpubs_inner_cdx.txt', (err3 or '') + '\n' + '\n'.join(' '.join(r) for r in rows))
        rows, err3 = cdx('rpubs.com/ElectProject/early_vote_2022', '20221001', '20221110')
        save('rpubs_cdx.txt', (err3 or '') + '\n' + '\n'.join(' '.join(r) for r in rows))

    # 5. Pennsylvania: every archived file under pavoterservices.pa.gov with "ballot" in the name, 2022-2024
    q = ('https://web.archive.org/cdx/search/cdx?url=pavoterservices.pa.gov/*&output=json&fl=timestamp,original,statuscode'
         '&from=2022&to=2024&filter=original:(?i).*ballot.*&collapse=original&limit=500')
    t, err = tryget(q, timeout=180)
    save('pa_cdx_all.txt', err or t)

    # 6. Nevada with a browser-like user agent
    for k, u in {'nv_turnout': 'https://www.nvsos.gov/sos/elections/election-information/election-statistics',
                 'nv_home': 'https://www.nvsos.gov/sos/'}.items():
        try:
            req = urllib.request.Request(u, headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/129.0 Safari/537.36',
                                                     'Accept': 'text/html,application/xhtml+xml', 'Accept-Language': 'en-US,en;q=0.9'})
            with urllib.request.urlopen(req, timeout=60) as r:
                html = r.read().decode('utf-8', errors='replace')
            save(f'{k}.txt', u + '\n' + text_preview(html, 3000) + '\n\nLINKS\n' + '\n'.join(f'{a} :: {tt}' for a, tt in links(html, u) if re.search(r'turnout|early|mail|statist', a + tt, re.I))[:8000])
        except Exception as e:
            save(f'{k}_error.txt', str(e))
    save('_log2.txt', '\n'.join(log))


if __name__ == '__main__':
    if sys.argv[1:2] == ['explore']:
        explore(sys.argv[2] if len(sys.argv) > 2 else 'ev_explore')
    elif sys.argv[1:2] == ['explore2']:
        explore2(sys.argv[2] if len(sys.argv) > 2 else 'ev_explore')
    else:
        print(__doc__)
